"""测试报告生成 —— HTML 与 Excel 双格式。

除了常规的用例结果与追溯矩阵，网络工程还多两块内容，报告里单列：

  * **可达性矩阵摘要**：按「源域 → 目的域」汇总放通/阻断数量。
    这是网络设计意图的量化快照；隔离要求如果被破坏，这里成片变色。
  * **故障注入结果**：每条注入故障是否被检出、症状是什么。
    用于验证校验能力的检出有效性（等价于变异测试）。

HTML 自包含（内联 CSS，无外部依赖），可直接当附件或贴进 PPT/文档。
"""

from __future__ import annotations

import html
import json
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional

from .traceability import (
    CaseMeta,
    CheckItem,
    RequirementMeta,
    audit_traceability,
    build_matrix,
    check_coverage,
    load_cases,
    load_checks,
    load_requirements,
)

STATUS_LABEL = {
    "passed": "通过", "failed": "失败", "skipped": "跳过",
    "not_run": "未执行", "未覆盖": "未覆盖", "部分执行": "部分执行",
    "通过": "通过", "失败": "失败",
}

STATUS_STYLE = {
    "passed": ("#E1F5EE", "#0F6E56", "#085041"),
    "failed": ("#FCEBEB", "#A32D2D", "#501313"),
    "skipped": ("#FAEEDA", "#854F0B", "#412402"),
    "not_run": ("#F1EFE8", "#5F5E5A", "#2C2C2A"),
    "未覆盖": ("#FAECE7", "#993C1D", "#4A1B0C"),
    "部分执行": ("#FAEEDA", "#854F0B", "#412402"),
    "通过": ("#E1F5EE", "#0F6E56", "#085041"),
    "失败": ("#FCEBEB", "#A32D2D", "#501313"),
}

KIND_STYLE = {
    "正向": ("#E6F1FB", "#185FA5"),
    "负向": ("#FAEEDA", "#854F0B"),
    "变异": ("#EEEDFE", "#534AB7"),
}

MONO = "font-family:ui-monospace,Consolas,'Courier New',monospace"


def _badge(status: str) -> str:
    bg, border, fg = STATUS_STYLE.get(status, STATUS_STYLE["not_run"])
    return (f'<span style="display:inline-block;padding:2px 10px;border-radius:10px;'
            f'background:{bg};border:1px solid {border};color:{fg};font-size:12px;'
            f'white-space:nowrap">{html.escape(STATUS_LABEL.get(status, status))}</span>')


def _kind_chips(kinds: Dict[str, int]) -> str:
    if not kinds:
        return '<span style="color:#B4B2A9">—</span>'
    out = ""
    for kind, count in sorted(kinds.items()):
        bg, fg = KIND_STYLE.get(kind, ("#F1EFE8", "#5F5E5A"))
        out += (f'<span style="display:inline-block;padding:1px 7px;border-radius:8px;'
                f'background:{bg};color:{fg};font-size:11px;margin-right:3px">'
                f'{html.escape(kind)} {count}</span>')
    return out


# --------------------------------------------------------------------------- #
@dataclass
class ReportData:
    project: str
    generated_at: str
    requirements: Dict[str, RequirementMeta]
    cases: Dict[str, CaseMeta]
    checks: List[CheckItem]
    results: Dict[str, dict]
    matrix: List[dict]
    audit: Dict[str, Any]
    coverage: Dict[str, Any]
    state: Dict[str, Any] = field(default_factory=dict)

    @property
    def total(self) -> int:
        return len(self.cases)

    @property
    def passed(self) -> int:
        return sum(1 for c in self.cases if self.results.get(c, {}).get("status") == "passed")

    @property
    def failed(self) -> int:
        return sum(1 for c in self.cases if self.results.get(c, {}).get("status") == "failed")

    @property
    def not_run(self) -> int:
        return sum(1 for c in self.cases if c not in self.results)

    @property
    def pass_rate(self) -> float:
        executed = self.total - self.not_run
        return self.passed / executed if executed else 0.0

    @property
    def duration(self) -> float:
        return sum(r.get("duration", 0.0) for r in self.results.values())

    def by_module(self) -> List[dict]:
        buckets: Dict[str, dict] = {}
        for cid, meta in self.cases.items():
            key = meta.module or meta.category or "未分类"
            b = buckets.setdefault(key, {"module": key, "total": 0, "passed": 0,
                                         "failed": 0, "not_run": 0, "kinds": {}})
            b["total"] += 1
            b["kinds"][meta.kind] = b["kinds"].get(meta.kind, 0) + 1
            status = self.results.get(cid, {}).get("status")
            if status == "passed":
                b["passed"] += 1
            elif status == "failed":
                b["failed"] += 1
            else:
                b["not_run"] += 1
        for b in buckets.values():
            executed = b["total"] - b["not_run"]
            b["pass_rate"] = b["passed"] / executed if executed else 0.0
        return sorted(buckets.values(), key=lambda x: x["module"])

    def failures(self) -> List[tuple]:
        return [(cid, self.cases[cid], self.results[cid])
                for cid in sorted(self.cases)
                if self.results.get(cid, {}).get("status") == "failed"]


# --------------------------------------------------------------------------- #
def load_report_data(
    requirements_path: str,
    cases_path: str,
    checks_path: str,
    results_path: str,
    state_path: Optional[str] = None,
    project: Optional[str] = None,
) -> ReportData:
    requirements = load_requirements(requirements_path)
    cases = load_cases(cases_path)
    checks = load_checks(checks_path)

    raw = json.loads(Path(results_path).read_text(encoding="utf-8"))
    results = raw.get("results", {})

    state: Dict[str, Any] = {}
    if state_path and Path(state_path).exists():
        state = json.loads(Path(state_path).read_text(encoding="utf-8"))

    return ReportData(
        project=project or "智能工厂网络 · 配置即代码与可达性验证",
        generated_at=raw.get("generated_at") or time.strftime("%Y-%m-%d %H:%M:%S"),
        requirements=requirements,
        cases=cases,
        checks=checks,
        results=results,
        matrix=build_matrix(requirements, cases, results),
        audit=audit_traceability(requirements, cases, results),
        coverage=check_coverage(checks, cases, results),
        state=state,
    )


# --------------------------------------------------------------------------- #
def render_html(data: ReportData) -> str:
    topo = data.state.get("topology", {})
    matrix = data.state.get("matrix", {})
    faults = data.state.get("faults", [])
    configs = data.state.get("configs", {})
    dead_rules = data.state.get("dead_rules", [])

    kpi = [
        ("用例总数", str(data.total), "#0C447C", "#E6F1FB", "#185FA5"),
        ("通过", str(data.passed), "#085041", "#E1F5EE", "#0F6E56"),
        ("失败", str(data.failed), "#501313", "#FCEBEB", "#A32D2D"),
        ("校验能力覆盖", f"{data.coverage['rate'] * 100:.0f}%", "#3C3489", "#EEEDFE", "#534AB7"),
        ("需求覆盖", f"{data.audit['coverage_rate'] * 100:.0f}%", "#633806", "#FAEEDA", "#854F0B"),
    ]
    kpi_html = "".join(
        f'<div style="flex:1;min-width:126px;background:{bg};border:1px solid {bd};'
        f'border-radius:12px;padding:14px 16px">'
        f'<div style="font-size:12px;color:{fg};opacity:.78">{html.escape(label)}</div>'
        f'<div style="font-size:25px;font-weight:500;color:{fg};margin-top:5px">{html.escape(value)}</div>'
        f"</div>"
        for label, value, fg, bg, bd in kpi
    )

    if topo:
        topo_html = (
            '<div style="display:flex;flex-wrap:wrap;gap:10px;font-size:13px">'
            + "".join(
                f'<div style="background:#FFFFFF;border:1px solid #D3D1C7;border-radius:10px;'
                f'padding:9px 14px"><span style="color:#5F5E5A">{html.escape(label)}</span> '
                f'<b style="color:#0C447C">{html.escape(str(value))}</b></div>'
                for label, value in [
                    ("安全域", topo.get("zones")), ("VLAN", topo.get("vlans")),
                    ("网络设备", topo.get("devices")), ("端系统", topo.get("hosts")),
                    ("服务类型", topo.get("services")), ("安全策略", topo.get("policies")),
                    ("链路", topo.get("links")), ("VRRP 组", topo.get("vrrp_groups")),
                ]
            )
            + "</div>"
        )
    else:
        topo_html = ""

    module_rows = "".join(
        f"<tr><td>{html.escape(b['module'])}</td>"
        f"<td style='text-align:right'>{b['total']}</td>"
        f"<td>{_kind_chips(b['kinds'])}</td>"
        f"<td style='text-align:right'>{b['passed']}</td>"
        f"<td style='text-align:right'>{b['failed']}</td>"
        f"<td style='text-align:right'>{b['pass_rate'] * 100:.1f}%</td></tr>"
        for b in data.by_module()
    )

    check_rows = "".join(
        f"<tr><td style='{MONO};font-size:12px;white-space:nowrap'>{html.escape(r['check'])}</td>"
        f"<td>{html.escape(r['name'])}</td>"
        f"<td>{html.escape(r['detects'])}</td>"
        f"<td>{_kind_chips(r['kinds'])}</td>"
        f"<td style='text-align:right'>{r['case_count']}</td>"
        f"<td>{_badge(r['status'])}</td></tr>"
        for r in data.coverage["rows"]
    )

    matrix_rows = "".join(
        f"<tr><td style='{MONO};font-size:12px'>{html.escape(r['requirement'])}</td>"
        f"<td>{html.escape(r['requirement_title'])}</td>"
        f"<td>{html.escape(r['module'])}</td>"
        f"<td style='{MONO};font-size:12px'>{html.escape('、'.join(r['cases']) or '—')}</td>"
        f"<td style='text-align:right'>{r['case_count']}</td>"
        f"<td>{_badge(r['status'])}</td></tr>"
        for r in data.matrix
    )

    # 可达性矩阵：按域对汇总
    zone_rows = ""
    zone_matrix = matrix.get("zone_matrix", {}) if matrix else {}
    for pair, counts in sorted(zone_matrix.items()):
        permit = counts.get("permit", 0)
        deny = counts.get("deny", 0)
        total = permit + deny
        ratio = (permit / total * 100) if total else 0
        bar_color = "#0F6E56" if permit == 0 else ("#A32D2D" if deny == 0 else "#854F0B")
        zone_rows += (
            f"<tr><td style='{MONO};font-size:12px'>{html.escape(pair)}</td>"
            f"<td style='text-align:right;color:#085041'>{permit}</td>"
            f"<td style='text-align:right;color:#A32D2D'>{deny}</td>"
            f"<td style='width:180px'><div style='height:8px;background:#F1EFE8;border-radius:4px'>"
            f"<div style='height:8px;width:{ratio:.0f}%;background:{bar_color};border-radius:4px'></div>"
            f"</div></td>"
            f"<td style='text-align:right'>{ratio:.0f}%</td></tr>"
        )

    fault_rows = ""
    for item in faults:
        ok = item.get("detected") == item.get("expect_detected")
        fault_rows += (
            f"<tr><td style='{MONO};font-size:12px'>{html.escape(item.get('fault_id', ''))}</td>"
            f"<td>{html.escape(item.get('name', ''))}</td>"
            f"<td>{_badge('通过' if ok else '失败')}</td>"
            f"<td>{'检出' if item.get('detected') else '未检出'}</td>"
            f"<td style='font-size:12px;color:#5F5E5A'>{html.escape(str(item.get('symptom', ''))[:90])}</td></tr>"
        )

    dead_rows = ""
    for item in dead_rules:
        dead_rows += (
            f"<tr><td style='{MONO};font-size:12px'>{html.escape(str(item.get('acl', '')))}</td>"
            f"<td style='{MONO};font-size:12px'>{html.escape(str(item.get('rule', '')))}</td>"
            f"<td style='font-size:12px'>{html.escape(str(item.get('reason', '')))}</td></tr>"
        )

    detail_rows = ""
    for cid in sorted(data.cases):
        meta = data.cases[cid]
        res = data.results.get(cid, {})
        status = res.get("status", "not_run")
        dur = f"{res.get('duration', 0):.2f}s" if res.get("duration") is not None else "—"
        detail_rows += (
            f"<tr><td style='{MONO};font-size:12px;white-space:nowrap'>{html.escape(cid)}</td>"
            f"<td>{html.escape(meta.title)}</td>"
            f"<td>{html.escape(meta.kind)}</td>"
            f"<td style='{MONO};font-size:12px'>{html.escape(meta.requirement)}</td>"
            f"<td>{html.escape(meta.method)}</td>"
            f"<td style='text-align:center'>{html.escape(meta.priority)}</td>"
            f"<td>{_badge(status)}</td>"
            f"<td style='text-align:right'>{dur}</td></tr>"
        )

    failures = data.failures()
    if failures:
        items = ""
        for cid, meta, res in failures:
            items += (
                f'<div style="border:1px solid #F7C1C1;background:#FCEBEB;border-radius:12px;'
                f'padding:14px 16px;margin-bottom:12px">'
                f'<div style="font-weight:500;color:#501313;margin-bottom:4px">'
                f"{html.escape(cid)} · {html.escape(meta.title)}</div>"
                f'<div style="font-size:12px;color:#A32D2D;margin-bottom:8px">'
                f"需求 {html.escape(meta.requirement)} ｜ 校验能力 "
                f"{html.escape('、'.join(meta.checks) or '—')} ｜ 节点 "
                f"{html.escape(res.get('nodeid', ''))}</div>"
                f'<pre style="margin:0;white-space:pre-wrap;font-size:12px;color:#2C2C2A;'
                f'{MONO}">{html.escape(res.get("error", "（无堆栈）"))}</pre></div>'
            )
        failure_html = f"<h2>失败用例详情（{len(failures)}）</h2>{items}"
    else:
        failure_html = (
            '<h2>失败用例详情</h2><div style="border:1px solid #9FE1CB;background:#E1F5EE;'
            'border-radius:12px;padding:16px;color:#085041">全部用例通过，无失败项。</div>'
        )

    audit = data.audit
    gaps: List[str] = []
    if audit["uncovered_requirements"]:
        gaps.append(f"未被用例覆盖的需求：{'、'.join(audit['uncovered_requirements'])}")
    if audit["dangling_requirement_refs"]:
        gaps.append(f"用例引用了不存在的需求：{'、'.join(audit['dangling_requirement_refs'])}")
    if audit["cases_not_executed"]:
        gaps.append(f"声明了但未执行的用例：{'、'.join(audit['cases_not_executed'])}")
    if audit["results_without_case_meta"]:
        gaps.append(f"执行结果未在用例清单登记：{'、'.join(audit['results_without_case_meta'])}")
    if data.coverage["uncovered"]:
        gaps.append(f"未被任何用例触发的校验能力：{'、'.join(data.coverage['uncovered'])}")
    gaps.extend(data.state.get("audit_extra", []))
    audit_html = ("".join(f"<li>{html.escape(g)}</li>" for g in gaps) if gaps
                  else "<li>需求、用例、校验能力与执行结果四方一致，无缺口。</li>")

    cfg_line = ""
    if configs:
        cfg_line = (f"生成配置 {configs.get('devices')} 台 / {configs.get('lines')} 行 / "
                    f"{configs.get('bytes')} 字节")

    return f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<title>{html.escape(data.project)} · 测试报告</title>
<style>
  body {{ margin:0; padding:32px 28px 56px; background:#FAFAF8;
         font-family:"Microsoft YaHei","PingFang SC",system-ui,sans-serif;
         color:#2C2C2A; line-height:1.65; }}
  .wrap {{ max-width:1120px; margin:0 auto; }}
  h1 {{ font-size:22px; font-weight:500; margin:0 0 6px; }}
  h2 {{ font-size:16px; font-weight:500; margin:32px 0 12px; padding-left:10px;
        border-left:3px solid #185FA5; }}
  .meta {{ color:#5F5E5A; font-size:13px; margin-bottom:22px; }}
  .kpis {{ display:flex; gap:12px; flex-wrap:wrap; margin-bottom:16px; }}
  table {{ width:100%; border-collapse:collapse; font-size:13px; background:#FFFFFF;
           border:1px solid #D3D1C7; border-radius:12px; overflow:hidden; }}
  th {{ background:#F1EFE8; text-align:left; padding:9px 12px; font-weight:500;
        color:#444441; border-bottom:1px solid #D3D1C7; white-space:nowrap; }}
  td {{ padding:8px 12px; border-bottom:1px solid #EDEBE4; vertical-align:top; }}
  tr:last-child td {{ border-bottom:none; }}
  ul {{ margin:0; padding-left:20px; font-size:13px; }}
  .foot {{ margin-top:32px; color:#888780; font-size:12px; }}
</style>
</head>
<body>
<div class="wrap">
  <h1>{html.escape(data.project)}</h1>
  <div class="meta">
    网络配置即代码与可达性验证报告 ｜ 生成时间 {html.escape(data.generated_at)}
    {f"｜ {html.escape(cfg_line)}" if cfg_line else ""}
    ｜ 累计用例耗时 {data.duration:.1f}s
  </div>

  <div class="kpis">{kpi_html}</div>

  {f'<h2>拓扑概览</h2>{topo_html}' if topo_html else ''}

  <h2>按模块汇总</h2>
  <table>
    <thead><tr><th>模块</th><th style="text-align:right">用例数</th><th>用例类型</th>
    <th style="text-align:right">通过</th><th style="text-align:right">失败</th>
    <th style="text-align:right">通过率</th></tr></thead>
    <tbody>{module_rows}</tbody>
  </table>

  <h2>校验能力清单（{data.coverage['covered']}/{data.coverage['total']}，
      {data.coverage['rate'] * 100:.1f}%）</h2>
  <table>
    <thead><tr><th>能力编号</th><th>能力名称</th><th>能查出什么</th>
    <th>用例类型</th><th style="text-align:right">用例数</th><th>状态</th></tr></thead>
    <tbody>{check_rows}</tbody>
  </table>

  <h2>需求追溯矩阵（RTM）</h2>
  <table>
    <thead><tr><th>需求编号</th><th>需求名称</th><th>模块</th><th>覆盖用例</th>
    <th style="text-align:right">用例数</th><th>状态</th></tr></thead>
    <tbody>{matrix_rows}</tbody>
  </table>

  <h2>可达性矩阵摘要（按源域 → 目的域）</h2>
  <div style="font-size:12px;color:#5F5E5A;margin-bottom:8px">
    全矩阵共 {matrix.get('total', 0)} 个流量组合，放通 {matrix.get('permitted', 0)}，
    阻断 {matrix.get('denied', 0)}；跨域组合 {matrix.get('cross_zone_total', 0)} 个。
    放通率越接近 0 说明隔离越严；某一格突然出现放通，通常意味着策略被误改。
  </div>
  <table>
    <thead><tr><th>源域 → 目的域</th><th style="text-align:right">放通</th>
    <th style="text-align:right">阻断</th><th>放通占比</th><th style="text-align:right">比例</th></tr></thead>
    <tbody>{zone_rows}</tbody>
  </table>

  <h2>故障注入结果（{sum(1 for f in faults if f.get('detected') == f.get('expect_detected'))}/{len(faults)}）</h2>
  <div style="font-size:12px;color:#5F5E5A;margin-bottom:8px">
    故意把拓扑改坏，要求校验套件必须报错。注入后依然全绿，说明校验有盲区。
  </div>
  <table>
    <thead><tr><th>编号</th><th>注入的故障</th><th>结论</th><th>是否检出</th><th>观察到的症状</th></tr></thead>
    <tbody>{fault_rows or '<tr><td colspan="5">无故障注入结果</td></tr>'}</tbody>
  </table>

  {f'<h2>ACL 死规则（{len(dead_rules)} 条）</h2><div style="font-size:12px;color:#5F5E5A;margin-bottom:8px">被前面的规则完全覆盖、永远不会命中的规则。不造成功能故障，但会误导排障。</div><table><thead><tr><th>ACL</th><th>规则</th><th>被谁覆盖</th></tr></thead><tbody>{dead_rows}</tbody></table>' if dead_rules else ''}

  <h2>追溯一致性审计</h2>
  <ul>{audit_html}</ul>

  {failure_html}

  <h2>全部用例明细</h2>
  <table>
    <thead><tr><th>用例编号</th><th>用例名称</th><th>类型</th><th>关联需求</th>
    <th>设计方法</th><th style="text-align:center">优先级</th>
    <th>结果</th><th style="text-align:right">耗时</th></tr></thead>
    <tbody>{detail_rows}</tbody>
  </table>

  <div class="foot">
    由 tools/make_report.py 自动生成 ｜ 数据来源 artifacts/results.json、
    artifacts/topology_state.json 与 cases/*.yaml<br>
    复现方式：<code>python tools/make_report.py --run</code>
  </div>
</div>
</body>
</html>
"""


# --------------------------------------------------------------------------- #
def write_html(data: ReportData, path: str) -> str:
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(render_html(data), encoding="utf-8")
    return str(out)


def write_excel(data: ReportData, path: str) -> Optional[str]:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return None

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="F1EFE8")
    head_font = Font(bold=True, color="444441")
    pass_fill = PatternFill("solid", fgColor="E1F5EE")
    fail_fill = PatternFill("solid", fgColor="FCEBEB")

    def header(ws, columns, widths=None):
        for i, title in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=i, value=title)
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(vertical="center")
        for i, title in enumerate(columns, start=1):
            ws.column_dimensions[get_column_letter(i)].width = (
                widths[i - 1] if widths else max(12, min(52, len(str(title)) * 2.4 + 10))
            )
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "概览"
    header(ws, ["指标", "数值"], [24, 64])
    topo = data.state.get("topology", {})
    for key, value in [
        ("项目", data.project), ("生成时间", data.generated_at),
        ("用例总数", data.total), ("通过", data.passed), ("失败", data.failed),
        ("未执行", data.not_run), ("通过率", f"{data.pass_rate * 100:.1f}%"),
        ("校验能力总数", data.coverage["total"]),
        ("校验能力已覆盖", data.coverage["covered"]),
        ("校验能力覆盖率", f"{data.coverage['rate'] * 100:.1f}%"),
        ("需求总数", data.audit["requirements_total"]),
        ("需求已覆盖", data.audit["requirements_covered"]),
        ("需求覆盖率", f"{data.audit['coverage_rate'] * 100:.1f}%"),
        ("累计耗时(s)", round(data.duration, 1)),
        ("—", ""),
        ("安全域", topo.get("zones")), ("VLAN", topo.get("vlans")),
        ("网络设备", topo.get("devices")), ("端系统", topo.get("hosts")),
        ("安全策略", topo.get("policies")), ("链路", topo.get("links")),
    ]:
        ws.append([key, value])

    ws = wb.create_sheet("用例明细")
    cols = ["用例编号", "用例名称", "模块", "类型", "关联需求", "触发的校验能力",
            "设计方法", "优先级", "结果", "耗时(s)", "节点"]
    header(ws, cols)
    for cid in sorted(data.cases):
        m = data.cases[cid]
        r = data.results.get(cid, {})
        status = r.get("status", "not_run")
        ws.append([cid, m.title, m.module, m.kind, m.requirement,
                   "、".join(m.checks), m.method, m.priority,
                   STATUS_LABEL.get(status, status), r.get("duration"),
                   r.get("nodeid", "")])
        cell = ws.cell(row=ws.max_row, column=9)
        if status == "failed":
            cell.fill = fail_fill
        elif status == "passed":
            cell.fill = pass_fill

    ws = wb.create_sheet("校验能力清单")
    header(ws, ["能力编号", "能力名称", "模块", "能查出什么", "说明", "用例数", "覆盖用例", "状态"],
           [16, 22, 14, 36, 40, 9, 34, 10])
    for r in data.coverage["rows"]:
        ws.append([r["check"], r["name"], r["module"], r["detects"],
                   r["description"], r["case_count"], "、".join(r["cases"]), r["status"]])

    ws = wb.create_sheet("需求追溯矩阵")
    header(ws, ["需求编号", "需求名称", "模块", "覆盖用例", "用例数", "状态", "需求出处"],
           [16, 30, 14, 34, 9, 10, 46])
    for row in data.matrix:
        ws.append([row["requirement"], row["requirement_title"], row["module"],
                   "、".join(row["cases"]), row["case_count"], row["status"], row["source"]])

    ws = wb.create_sheet("可达性矩阵")
    header(ws, ["源域 → 目的域", "放通", "阻断", "放通占比"], [26, 10, 10, 12])
    for pair, counts in sorted(data.state.get("matrix", {}).get("zone_matrix", {}).items()):
        permit = counts.get("permit", 0)
        deny = counts.get("deny", 0)
        total = permit + deny
        ws.append([pair, permit, deny, f"{permit / total * 100:.1f}%" if total else "0%"])

    ws = wb.create_sheet("故障注入")
    header(ws, ["编号", "注入的故障", "变异算子", "期望检出", "实际检出", "结论", "症状"],
           [12, 34, 26, 10, 10, 10, 60])
    for item in data.state.get("faults", []):
        ok = item.get("detected") == item.get("expect_detected")
        ws.append([item.get("fault_id"), item.get("name"), item.get("mutation"),
                   item.get("expect_detected"), item.get("detected"),
                   "通过" if ok else "失败", item.get("symptom", "")])

    ws = wb.create_sheet("失败详情")
    header(ws, ["用例编号", "用例名称", "关联需求", "触发的校验能力", "失败节点", "错误信息"],
           [16, 30, 16, 30, 40, 70])
    for cid, meta, res in data.failures():
        ws.append([cid, meta.title, meta.requirement, "、".join(meta.checks),
                   res.get("nodeid", ""), res.get("error", "")])
    if not data.failures():
        ws.append(["—", "无失败用例", "", "", "", ""])

    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    return str(out)


def render_console_summary(data: ReportData) -> str:
    matrix = data.state.get("matrix", {})
    faults = data.state.get("faults", [])
    lines = [
        "=" * 78,
        f"  {data.project}",
        "=" * 78,
        f"  生成时间        {data.generated_at}",
        f"  用例总数        {data.total}",
        f"  通过 / 失败     {data.passed} / {data.failed}",
        f"  通过率          {data.pass_rate * 100:.1f}%",
        f"  校验能力覆盖率  {data.coverage['rate'] * 100:.1f}%"
        f"（{data.coverage['covered']}/{data.coverage['total']}）",
        f"  需求覆盖率      {data.audit['coverage_rate'] * 100:.1f}%"
        f"（{data.audit['requirements_covered']}/{data.audit['requirements_total']}）",
    ]
    if matrix:
        lines.append(f"  可达性矩阵      {matrix.get('total')} 个组合，放通 "
                     f"{matrix.get('permitted')}，阻断 {matrix.get('denied')}")
    if faults:
        ok = sum(1 for f in faults if f.get("detected") == f.get("expect_detected"))
        lines.append(f"  故障注入检出    {ok}/{len(faults)}")
    lines.append("-" * 78)
    for b in data.by_module():
        lines.append(
            f"  {b['module']:<20} 共 {b['total']:>2} 条  "
            f"通过 {b['passed']:>2}  失败 {b['failed']:>2}  "
            f"通过率 {b['pass_rate'] * 100:5.1f}%"
        )
    failures = data.failures()
    if failures:
        lines.append("-" * 78)
        lines.append("  失败用例：")
        for cid, meta, _ in failures:
            lines.append(f"    - {cid}  {meta.title}（{meta.requirement}）")
    gaps = (data.audit["uncovered_requirements"] + data.audit["dangling_requirement_refs"]
            + data.audit["cases_not_executed"] + data.coverage["uncovered"])
    gaps.extend(data.state.get("audit_extra", []))
    if gaps:
        lines.append("-" * 78)
        lines.append(f"  追溯缺口：{'、'.join(gaps)}")
    lines.append("=" * 78)
    return "\n".join(lines)
